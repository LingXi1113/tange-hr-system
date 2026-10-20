from datetime import date

from common.resume_parser import extract_text, parse_resume_fields, parse_resume_file_with_text


def test_resume_parser_extracts_age_highest_education_and_major():
    fields = parse_resume_fields("""
姓名：张三
性别：男
年龄：28岁
教育经历
2015.09-2019.06 华东理工大学 计算机科学与技术 本科
2019.09-2022.06 中山大学 软件工程 硕士研究生
""")

    assert fields["age"] == 28
    assert fields["highest_education"] == "硕士"
    assert fields["major"] == "软件工程"
    assert fields["education"][0]["degree"] == "本科"
    assert fields["education"][1]["degree"] == "硕士"


def test_resume_parser_derives_age_from_birth_date_and_labelled_major():
    fields = parse_resume_fields("""
姓名：李四
出生日期：1998-10-20
教育经历
2017-2021 广东工业大学
专业：市场营销
学历：本科
""")

    expected_age = date.today().year - 1998 - (
        (date.today().month, date.today().day) < (10, 20)
    )
    assert fields["age"] == expected_age
    assert fields["highest_education"] == "本科"
    assert fields["major"] == "市场营销"


def test_resume_parser_normalizes_pdf_radicals_and_separates_school_department_major():
    fields = parse_resume_fields("""
卢思源
男 21岁
中⼭⼤学东校园
教育背景
中⼭⼤学（985） 计算机学院/计算机科学与技术 本科 GPA: 3.3/4.0
2023-09 - 2027-06
""")

    assert fields["education"] == [{
        "school": "中山大学",
        "major": "计算机科学与技术",
        "degree": "本科",
        "graduate_at": "2027-06",
    }]
    assert fields["highest_education"] == "本科"
    assert fields["major"] == "计算机科学与技术"


def test_resume_parser_infers_vocational_college_degree_and_year_range():
    fields = parse_resume_fields("""
教育经历
海南职业技术学院 | 工程造价 | 2021-2024
资格证书
驾驶证 C1
""")

    assert fields["education"] == [{
        "school": "海南职业技术学院",
        "major": "工程造价",
        "degree": "大专",
        "graduate_at": "2024",
    }]
    assert fields["highest_education"] == "大专"
    assert fields["major"] == "工程造价"


def test_resume_parser_prefers_person_name_over_job_title_in_filename():
    fields = parse_resume_fields(
        "手机：19535885461\n邮箱：cedwang@163.com\n最高学历：硕士",
        "战略投资-王轶群.pdf",
    )

    assert fields["name"] == "王轶群"


def test_resume_parser_prefers_unlabelled_header_name_over_filename():
    fields = parse_resume_fields(
        "王轶群\n手机：19535885461\n邮箱：cedwang@163.com",
        "战略投资.pdf",
    )

    assert fields["name"] == "王轶群"


def test_resume_parser_scopes_and_merges_education_records():
    fields = parse_resume_fields("""
王轶群
个人简介
参与北京大学战略投资项目，岗位要求硕士优先
教育背景
2018-09 - 2020-06 伦敦商学院 International Business 硕士
2018-09 - 2020-06 伦敦商学院 International Business 硕士
工作经历
2020-07 - 至今
某某投资有限公司
战略投资经理
""")

    assert fields["name"] == "王轶群"
    assert fields["highest_education"] == "硕士"
    assert len(fields["education"]) == 1
    assert fields["education"][0]["school"] == "伦敦商学院"
    assert fields["education"][0]["major"] == "International Business"


def test_resume_parser_does_not_treat_faculty_as_school():
    fields = parse_resume_fields("""
教育经历
2020-09 - 2024-06
International Business 商学院
硕士
""")

    assert fields["education"] == []


def test_resume_parser_reads_docx_tables(tmp_path):
    import docx

    path = tmp_path / "table-resume.docx"
    document = docx.Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "王轶群"
    table.cell(1, 0).text = "手机"
    table.cell(1, 1).text = "19535885461"
    table.cell(2, 0).text = "邮箱"
    table.cell(2, 1).text = "cedwang@163.com"
    document.save(path)

    text = extract_text(str(path))
    assert "姓名 | 王轶群" in text
    assert "19535885461" in text
    fields = parse_resume_fields(text, path.name)
    assert fields["name"] == "王轶群"
    assert fields["phone"] == "19535885461"
    assert fields["email"] == "cedwang@163.com"


def test_scanned_pdf_uses_local_ocr(monkeypatch):
    monkeypatch.setattr("common.resume_parser._extract_pdf", lambda _path: "")
    monkeypatch.setattr(
        "common.resume_parser.OcrService.recognize",
        lambda _self, _path: ("姓名：王轶群\n手机：19535885461\n邮箱：cedwang@163.com", 0.93),
    )

    fields, status, text = parse_resume_file_with_text("scan.pdf", "战略投资-王轶群.pdf")

    assert status == "system"
    assert fields["name"] == "王轶群"
    assert fields["_parse_meta"]["method"] == "ocr"
    assert fields["_parse_meta"]["confidence"] <= 93
    assert "19535885461" in text


def test_low_confidence_result_uses_ai_fallback(monkeypatch):
    monkeypatch.setattr(
        "common.resume_parser.extract_text_with_method",
        lambda _path: ("王轶群\n教育背景\n伦敦商学院 International Business 硕士", "text", 1.0),
    )
    monkeypatch.setattr("common.resume_ai.is_configured", lambda: True)
    monkeypatch.setattr(
        "common.resume_ai.extract_resume_fields",
        lambda _text, _filename: ({
            "name": "王轶群",
            "highest_education": "硕士",
            "education": [{
                "school": "伦敦商学院", "major": "International Business",
                "degree": "硕士", "graduate_at": "2020-06",
            }],
            "work_experience": [],
        }, ""),
    )

    fields, status, _ = parse_resume_file_with_text("resume.pdf", "战略投资-王轶群.pdf")

    assert status == "system"
    assert fields["education"][0]["school"] == "伦敦商学院"
    assert fields["highest_education"] == "硕士"
    assert fields["_parse_meta"]["method"] == "text+ai"


def test_openai_compatible_ai_adapter(monkeypatch):
    import common.resume_ai as resume_ai

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self):
            return b'{"choices":[{"message":{"content":"{\\"name\\":\\"Wang Yiqun\\"}"}}]}'

    monkeypatch.setenv("HRATS_RESUME_AI_ENABLED", "1")
    monkeypatch.setenv("HRATS_RESUME_AI_MODEL", "local-test-model")
    monkeypatch.setenv("HRATS_RESUME_AI_BASE_URL", "http://127.0.0.1:11434/v1")
    monkeypatch.setattr(resume_ai, "urlopen", lambda _request, timeout: FakeResponse())

    fields, error = resume_ai.extract_resume_fields("Name: Wang Yiqun", "resume.pdf")

    assert error == ""
    assert fields["name"] == "Wang Yiqun"


def test_ai_fallback_does_not_invent_missing_education():
    from common.resume_parser import _merge_ai_fields

    fields = parse_resume_fields("姓名：赵明\n教育背景", "赵明.pdf")
    merged = _merge_ai_fields(fields, {
        "highest_education": "初中",
        "education": [{"school": "虚构大学", "degree": "本科"}],
        "major": "计算机科学",
        "work_experience": [{"company": "虚构公司"}],
    }, "姓名：赵明\n教育背景", "赵明.pdf")
    assert merged["education"] == []
    assert merged["work_experience"] == []
    assert not merged["highest_education"]
    assert not merged["major"]

