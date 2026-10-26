"""Resume text extraction and conservative structured parsing.

The parser follows the useful parts of orgatAI/resume-parse-python:
extract PDF text with a layout-aware parser first, then associate nearby
dates, schools, degrees and majors into education records. It is deliberately
conservative: results are drafts and must remain editable by HR.
"""
import logging
import os
import re
import unicodedata
from datetime import date

from common.resume_work import WORK_PERIOD_RE, extract_work_experience
from common.resume_work import normalize_work_date as _normalize_work_date

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp"}
PDF_EXTS = {".pdf"}
WORD_EXTS = {".docx", ".doc"}
logger = logging.getLogger(__name__)

DEGREE_LEVELS = {
    "初中": 1,
    "高中": 2,
    "中专": 2,
    "大专": 3,
    "本科": 4,
    "硕士": 5,
    "博士": 6,
}

DEGREE_RE = re.compile(
    r"博士后|博士研究生|博士在读|工商管理硕士|工程硕士|专业硕士|在职研究生|硕士研究生|研究生|本科|大专|高职|专科|中专|高中|初中|博士|硕士|学士"
    r"|\b(?:post[- ]?doc(?:toral)?|Ph\.?\s*D\.?|D\.?Phil\.?|Doctor(?:ate|\s+of\s+Philosophy)"
    r"|Master(?:['’]s)?(?:\s+of\s+(?:Science|Arts|Engineering|Business\s+Administration))?"
    r"|M\.?\s*(?:Sc|Eng|Phil|Res)\.?|M\.?A\.?|EMBA|MBA"
    r"|Bachelor(?:['’]s)?(?:\s+of\s+(?:Science|Arts|Engineering))?"
    r"|B\.?\s*(?:Sc|Eng)\.?|B\.?A\.?)(?![A-Za-z])",
    flags=re.IGNORECASE,
)
ENGLISH_SCHOOL_RE = re.compile(
    r"\b(?:University\s+of\s+[A-Z][A-Za-zÀ-ž'’.-]*(?:\s+(?:of|and|the|[A-Z][A-Za-zÀ-ž'’.-]*)){0,5}"
    r"|(?:[A-Z][A-Za-zÀ-ž'’.-]*\s+){1,7}(?:University|College|School)"
    r"|(?:[A-Z][A-Za-zÀ-ž'’.-]*\s+){1,5}Institute\s+of\s+Technology)\b",
)
SCHOOL_RE = re.compile(
    r"[\u4e00-\u9fffA-Za-z0-9·&.\-]{2,50}?(?:大学|学院|学校|分校|电大)",
)
DATE_RE = re.compile(
    r"(?:19|20)\d{2}(?:\s*(?:年|[./-])\s*(?:0?[1-9]|1[0-2])(?!\d)\s*月?)?|(?:至今|现在)",
)
MAJOR_LABEL_RE = re.compile(
    r"(?:专业|主修|研究方向|所学专业|专业方向)\s*[：:|｜]\s*([^\r\n]{2,40})",
    flags=re.IGNORECASE,
)
GENDER_LABEL_RE = re.compile(
    r"(?:\u6027\u522b|gender|sex)\s*[\uFF1A:|｜,\uFF0C\s]\s*(\u7537|\u5973|male|female|m|f)",
    flags=re.IGNORECASE,
)
DATE_TOKEN = r"(?:19|20)\d{2}\s*(?:[\u5E74./-]\s*\d{1,2}\s*\u6708?)?"
WORK_SECTION_HEADINGS = {
    "\u5DE5\u4F5C\u7ECF\u5386", "\u5DE5\u4F5C\u7ECF\u9A8C", "\u804C\u4E1A\u7ECF\u5386", "\u5DE5\u4F5C\u5C65\u5386", "\u4EFB\u804C\u7ECF\u5386", "employment",
    "work experience", "professional experience", "career history",
}
SECTION_HEADINGS = {
    "个人优势", "个人信息", "个人简介", "个人简历", "简历", "求职意向", "联系方式",
    "教育经历", "教育背景", "工作经历", "项目经历", "实习经历", "技能特长", "专业技能",
    "自我评价", "证书奖励", "荣誉奖项", "兴趣爱好", "语言能力", "校园经历", "基本信息",
    "resume", "curriculum vitae", "cv",
}
# Keep the section vocabulary Unicode-safe even when a legacy source file was
# checked out with a non-UTF-8 console encoding.
SECTION_HEADINGS = {
    "\u4E2A\u4EBA\u4F18\u52BF", "\u4E2A\u4EBA\u4FE1\u606F", "\u4E2A\u4EBA\u7B80\u4ECB", "\u4E2A\u4EBA\u7B80\u5386", "\u7B80\u5386",
    "\u6C42\u804C\u610F\u5411", "\u8054\u7CFB\u65B9\u5F0F", "\u6559\u80B2\u7ECF\u5386", "\u6559\u80B2\u80CC\u666F",
    "\u5DE5\u4F5C\u7ECF\u5386", "\u9879\u76EE\u7ECF\u5386", "\u5B9E\u4E60\u7ECF\u5386", "\u6280\u80FD\u7279\u957F", "\u4E13\u4E1A\u6280\u80FD",
    "\u81EA\u6211\u8BC4\u4EF7", "\u8BC1\u4E66\u5956\u52B1", "\u8363\u8A89\u5956\u9879", "\u5174\u8DA3\u7231\u597D", "\u8BED\u8A00\u80FD\u529B",
    "\u6821\u56ED\u7ECF\u5386", "\u57FA\u672C\u4FE1\u606F", "resume", "curriculum vitae", "cv",
}
ALL_SECTION_HEADINGS = SECTION_HEADINGS | WORK_SECTION_HEADINGS

EDUCATION_SECTION_HEADINGS = {
    "教育经历", "教育背景", "学历背景", "学习经历", "education", "education background",
}

# 文件名中经常把“岗位-姓名”写在一起。这些词应作为岗位/部门证据，不能当成人名。
NAME_BLOCKLIST_TERMS = {
    "投资", "战略", "销售", "运营", "市场", "商务", "产品", "技术", "研发", "开发",
    "工程师", "经理", "主管", "总监", "专员", "助理", "顾问", "实习", "应聘", "候选人",
    "简历", "招聘", "人力资源", "财务", "行政", "法务", "设计", "采购", "客服", "岗位",
}
COMMON_CHINESE_SURNAMES = set(
    "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜戚谢邹喻柏水窦章云苏潘葛奚范彭郎鲁韦昌马苗凤花方俞任袁柳鲍史唐费廉岑薛雷贺倪汤滕殷罗毕郝邬安常乐于时傅皮卞齐康伍余元卜顾孟平黄和穆萧尹姚邵湛汪祁毛禹狄米贝明臧计伏成戴宋茅庞熊纪舒屈项祝董梁杜阮蓝闵席季麻强贾路娄危江童颜郭梅盛林刁钟徐邱骆高夏蔡田樊胡凌霍虞万支柯管卢莫经房裘缪干解应宗丁宣邓郁单杭洪包诸左石崔吉龚程嵇邢滑裴陆荣翁荀羊甄曲封芮储靳汲邴糜松井段富巫乌焦巴弓牧隗山谷车侯宓蓬全班仰秋仲伊宫宁仇栾暴甘钭厉戎祖武符刘景詹束龙叶幸司韶郜黎蓟薄印宿白怀蒲邰从鄂索咸籍赖卓蔺屠蒙池乔阴郁胥能苍双闻莘党翟谭贡劳逄姬申扶堵冉宰郦雍桑桂濮牛寿通边扈燕冀浦尚农温别庄晏柴瞿阎慕连茹习宦艾鱼容向古易慎戈廖庾终暨居衡步都耿满弘匡国文寇广禄阙东欧殳沃利蔚越夔隆师巩厍聂晁勾敖融冷訾辛阚那简饶空曾毋沙乜养鞠须丰巢关蒯相查后荆红游竺权逯盖益桓公"
)

# 这些“学院”一般是大学内部院系；缺少上级大学时宁可留空，也不要伪造院校。
FACULTY_ONLY_NAMES = {
    "商学院", "管理学院", "经济学院", "金融学院", "法学院", "外国语学院", "计算机学院",
    "软件学院", "人工智能学院", "新闻学院", "传播学院", "设计学院", "国际商学院",
}


class OcrService:
    """Lazy local OCR adapter; import failures degrade to manual review."""

    def recognize(self, file_path: str):  # noqa: D102
        from common.resume_ocr import recognize_resume

        return recognize_resume(file_path)


def _pdf_rows(layout) -> list[str]:
    """Restore visual rows before PDFMiner's column reading order loses links."""
    from pdfminer.layout import LTTextLine

    fragments = []

    def visit(item):
        if isinstance(item, LTTextLine):
            value = re.sub(r"\s+", " ", item.get_text()).strip()
            if value:
                fragments.append((item.x0, (item.y0 + item.y1) / 2, max(item.height, 1), value))
        elif hasattr(item, "__iter__"):
            for child in item:
                visit(child)

    visit(layout)
    rows = []
    for fragment in sorted(fragments, key=lambda item: (-item[1], item[0])):
        if rows and abs(rows[-1][0] - fragment[1]) <= min(rows[-1][1], fragment[2]) * 0.35:
            rows[-1][2].append(fragment)
        else:
            rows.append([fragment[1], fragment[2], [fragment]])
    return [" | ".join(item[3] for item in sorted(row[2], key=lambda item: item[0])) for row in rows]


def _extract_pdf(file_path: str) -> str:
    """Extract PDF text with layout analysis, falling back to pypdf."""
    text = ""
    try:
        from pdfminer.high_level import extract_pages
        from pdfminer.layout import LAParams

        text = "\n".join(
            line for page in extract_pages(
                file_path,
                laparams=LAParams(char_margin=2.0, line_margin=0.4, word_margin=0.1, boxes_flow=None),
            ) for line in _pdf_rows(page)
        )
    except Exception:
        logger.exception("PDFMiner 文本提取失败 extension=pdf")

    if text.strip():
        return text

    try:
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        logger.exception("pypdf 文本提取失败 extension=pdf")
        return ""


def _extract_docx(file_path: str) -> str:
    import docx

    document = docx.Document(file_path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # 很多招聘网站导出的 DOCX 把基础信息、教育和工作经历放在表格中。
    for table in document.tables:
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _text_quality(text: str) -> int:
    """Count useful characters; scanned PDFs usually contain almost none."""
    return len(re.findall(r"[\u4e00-\u9fffA-Za-z0-9@]", str(text or "")))


def extract_text_with_method(file_path: str) -> tuple[str, str, float]:
    """Return text, extraction method and OCR confidence."""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in PDF_EXTS:
            native_text = _extract_pdf(file_path) or ""
            threshold = max(20, int(os.environ.get("HRATS_RESUME_OCR_TEXT_THRESHOLD", "80")))
            if _text_quality(native_text) >= threshold:
                return native_text, "text", 1.0
            ocr_text, confidence = OcrService().recognize(file_path)
            if _text_quality(ocr_text) > _text_quality(native_text):
                return ocr_text, "ocr", confidence
            return native_text, "text", 1.0 if native_text.strip() else 0.0
        if ext == ".docx":
            return _extract_docx(file_path), "text", 1.0
        if ext in IMAGE_EXTS:
            text, confidence = OcrService().recognize(file_path)
            return text, "ocr", confidence
    except Exception:
        logger.exception("简历文本提取失败 extension=%s", ext)
        return "", "failed", 0.0
    return "", "failed", 0.0


def extract_text(file_path: str):
    """Backward-compatible text-only extraction entry point."""
    text, _, _ = extract_text_with_method(file_path)
    return text


def _normalize_date(value: str) -> str:
    value = re.sub(r"\s+", "", value)
    if value in {"至今", "现在"}:
        return "至今"
    matched = re.match(r"((?:19|20)\d{2})[年./-]?(\d{1,2})?", value)
    if not matched:
        return value
    year, month = matched.groups()
    return f"{year}-{int(month):02d}" if month else year


def _clean_school(value: str) -> str:
    value = re.sub(r"^(?:教育经历|教育背景|毕业院校|学校)\s*[：:，,]?", "", value)
    period = WORK_PERIOD_RE.match(value)
    if period:
        value = value[period.end():]
    else:
        value = re.sub(rf"^{DATE_TOKEN}\s*", "", value)
    return value.strip(" ：:，,;；|·")


def _normalize_resume_text(text: str) -> str:
    """Normalize visually identical compatibility glyphs emitted by some PDFs.

    For example, pdfminer may return ``中⼭⼤学`` where the mountain/big
    characters are Kangxi radicals instead of normal CJK characters.
    """
    value = unicodedata.normalize("NFKC", str(text or ""))
    return value.replace("\u00a0", " ").replace("\u200b", "")


def _school_candidates(line: str) -> list[str]:
    matches = []
    for value in SCHOOL_RE.findall(line):
        school = _clean_school(value)
        if school and school not in matches:
            matches.append(school)
    # A university and its internal faculty often occur on the same line.
    # In that case the university is the educational institution; the faculty
    # must not replace it (e.g. 中山大学 计算机学院/计算机科学与技术).
    universities = [
        item for item in matches
        if item.endswith(("大学", "学校", "分校", "电大"))
    ]
    if universities or matches:
        return universities or matches
    # Degree/major columns must not become part of an English institution name.
    prefix = DEGREE_RE.split(line, maxsplit=1)[0]
    return list(dict.fromkeys(match.strip() for match in ENGLISH_SCHOOL_RE.findall(prefix)))


def _pick_degree(text: str) -> str:
    text = re.split(r"(?:导师|指导教师)\s*[：:]", text, maxsplit=1)[0]
    text = re.sub(r"(?:博士|硕士)生?导师", "", text)
    matches = [normalize_degree(value) for value in DEGREE_RE.findall(text)]
    return max((value for value in matches if value), key=lambda value: DEGREE_LEVELS[value], default="")


def _pick_major(text: str) -> str:
    matched = MAJOR_LABEL_RE.search(text)
    if not matched:
        return ""
    return re.split(r"(?:学历|学位|学校|毕业时间)\s*[：:]", matched.group(1))[0].strip(" ：:，,;；|")


def normalize_degree(value: str) -> str:
    """Normalize common degree names for storage and reporting."""
    value = str(value or "").strip()
    if not value:
        return ""
    if re.search(r"博士(?!后)|\b(?:Ph\.?\s*D\.?|D\.?Phil\.?|Doctorate|Doctor\s+of\s+Philosophy)(?![A-Za-z])", value, re.I):
        return "博士"
    if re.search(r"硕士|\b(?:EMBA|MBA|Master(?:['’]s)?|M\.?\s*(?:Sc|Eng|Phil|Res)\.?|M\.?A\.?)(?![A-Za-z])", value, re.I):
        return "硕士"
    if re.search(r"本科|学士|\b(?:Bachelor(?:['’]s)?|B\.?\s*(?:Sc|Eng)\.?|B\.?A\.?)(?![A-Za-z])", value, re.I):
        return "本科"
    if any(token in value for token in ("大专", "专科", "高职")):
        return "大专"
    if "中专" in value:
        return "中专"
    if "高中" in value:
        return "高中"
    if "初中" in value:
        return "初中"
    return ""


def _infer_degree_from_school(school: str) -> str:
    """Infer only institution types whose education level is unambiguous enough."""
    school = str(school or "").strip()
    if any(token in school for token in ("职业技术学院", "职业学院", "高等专科学校", "专科学校")):
        return "大专"
    if any(token in school for token in ("中等专业学校", "中等职业学校")):
        return "中专"
    return ""


def _education_rank(record: dict) -> tuple[int, str]:
    degree = normalize_degree(record.get("degree", ""))
    return DEGREE_LEVELS.get(degree, 0), str(record.get("graduate_at") or "")


def highest_education_record(records: list[dict]) -> dict:
    return max(records, key=_education_rank, default={})


def highest_education_from_records(records: list[dict], text: str = "") -> str:
    """Return the highest normalized education level found in records/text."""
    candidates = [normalize_degree(item.get("degree", "")) for item in records]
    labelled = re.search(r"最高学历\s*[：:|｜]?\s*([^\r\n|｜]{1,30})", text or "")
    if labelled:
        candidates.append(_pick_degree(labelled.group(1)))
    # 已识别到教育记录时，只使用记录内学历，避免把“岗位要求硕士”等正文误作候选人学历。
    if not any(candidates):
        candidates.extend(_pick_degree(line) for line in _education_lines(text or "") if _education_degree_line(line))
    candidates = [item for item in candidates if item]
    return max(candidates, key=lambda item: DEGREE_LEVELS.get(item, 0), default="")


def candidate_highest_education(candidate: dict) -> str:
    """Use complete histories without discarding a manually supplied higher level."""
    return highest_education_from_records([
        *(candidate.get("education") or []),
        {"degree": candidate.get("highest_education", "")},
    ])


def major_from_records(records: list[dict], text: str = "") -> str:
    """Prefer the major attached to the highest (then latest) education record."""
    with_major = [item for item in records if str(item.get("major") or "").strip()]
    if with_major:
        selected = max(with_major, key=_education_rank)
        return str(selected.get("major") or "").strip()
    return _pick_major(text or "")


def _extract_age(text: str) -> int | None:
    """Extract an explicit age or derive it from a labelled birth date/year."""
    direct_patterns = (
        r"(?:年龄|age)\s*[：:]?\s*(\d{2})\s*(?:岁|周岁)?",
        r"(?<!\d)(\d{2})\s*岁(?!\d)",
    )
    for pattern in direct_patterns:
        matched = re.search(pattern, text, flags=re.IGNORECASE)
        if matched:
            value = int(matched.group(1))
            if 16 <= value <= 75:
                return value

    birth = re.search(
        r"(?:出生日期|出生年月|生日|出生)\s*[：:]?\s*((?:19|20)\d{2})"
        r"(?:\s*[年./-]\s*(\d{1,2}))?(?:\s*[月./-]\s*(\d{1,2}))?",
        text,
        flags=re.IGNORECASE,
    ) or re.search(r"((?:19|20)\d{2})\s*年\s*生", text)
    if not birth:
        return None
    year = int(birth.group(1))
    month = int(birth.group(2)) if birth.lastindex and birth.lastindex >= 2 and birth.group(2) else 1
    day = int(birth.group(3)) if birth.lastindex and birth.lastindex >= 3 and birth.group(3) else 1
    today = date.today()
    value = today.year - year - ((today.month, today.day) < (month, day))
    return value if 16 <= value <= 75 else None


def _major_from_school_line(line: str, school: str) -> str:
    """Handle compact rows such as '2020-2024 某大学 计算机 本科'."""
    value = line.replace(school, " ")
    value = re.sub(r"(?:工科|理科|文科|学历|学位)\s*[/／]\s*", " ", value)
    value = WORK_PERIOD_RE.sub(" ", value)
    value = re.sub(r"(?:19|20)\d{2}\s*(?:[年./-]\s*\d{1,2}\s*月?)?", " ", value)
    value = DEGREE_RE.sub(" ", value)
    value = re.sub(r"[（(]\s*(?:985|211|双一流|一本|二本)\s*[）)]", " ", value, flags=re.I)
    value = re.sub(r"\bGPA\s*[:：]?\s*[\d.]+(?:\s*/\s*[\d.]+)?", " ", value, flags=re.I)
    value = re.sub(r"(?:专业|主修|研究方向|所学专业|专业方向)\s*[：:]?", " ", value)
    value = re.sub(r"[\u4e00-\u9fffA-Za-z]{2,30}(?:学院|学部|系)(?=\s|[/|｜·，,;；]|$)", " ", value)
    parts = [
        re.sub(r"\s+", " ", item).strip(" ：:-~—–")
        for item in re.split(r"[/|｜·，,;；]+", value)
    ]
    candidates = [
        item for item in parts
        if 2 <= len(item) <= 30
        and not _is_section_heading(item)
        and "课程" not in item
        and not item.endswith(("校园", "校区"))
        and not item.upper().startswith("GPA")
    ]
    if candidates:
        return candidates[-1]
    return ""


def _normalized_heading(line: str) -> str:
    return re.sub(r"[：:\s]+$", "", str(line or "").strip()).lower()


def _education_lines(text: str) -> list[str]:
    """Prefer explicit education sections; fall back to all lines for simple resumes."""
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    sections: list[str] = []
    found_heading = False
    collecting = False
    for line in lines:
        parts = re.split(r"[|｜]", line, maxsplit=1)
        heading = _normalized_heading(parts[0])
        if heading in EDUCATION_SECTION_HEADINGS:
            found_heading = True
            collecting = True
            if len(parts) > 1 and parts[1].strip():
                sections.append(parts[1].strip())
            continue
        if collecting and (_is_section_heading(line) or _is_section_heading(parts[0])):
            collecting = False
            continue
        if collecting:
            sections.append(line)
    return sections if found_heading else lines


def _education_context(lines: list[str], index: int) -> str:
    """Build a small record context without leaking fields from adjacent schools."""
    selected = [lines[index]]
    for offset in (1, 2):
        previous = index - offset
        if previous < 0 or _real_schools(lines[previous]) or _is_section_heading(lines[previous]):
            break
        # Only borrow a preceding date row, not the previous school's degree.
        if DATE_RE.search(lines[previous]) and not _pick_degree(lines[previous]):
            selected.insert(0, lines[previous])
        else:
            break
    for offset in range(1, 13):
        following = index + offset
        if following >= len(lines) or _real_schools(lines[following]) or _is_section_heading(lines[following]):
            break
        if _education_degree_line(lines[following]):
            selected.append(lines[following])
    return "\n".join(selected)


def _real_schools(line: str) -> list[str]:
    return [school for school in _school_candidates(line) if school not in FACULTY_ONLY_NAMES]


def _education_degree_line(line: str) -> bool:
    """Reject coursework and JD requirements as evidence of an earned degree."""
    return not re.search(r"课程|论文题目|导师|指导教师|岗位要求|职位要求|招聘要求|学历要求|硕士优先|博士优先|\bcourses?\b|\brequirements?\b", line, re.I)


def _column_degrees(lines: list[str]) -> dict[int, str]:
    """Recover a detached degree column only when a one-to-one order exists.

    Native PDFs are fixed by coordinates first. OCR/legacy text can still put
    a full degree column after the institution column. Ambiguous counts are
    deliberately left unassigned for HR/AI review, never guessed by proximity.
    """
    schools = [(index, line) for index, line in enumerate(lines) if _real_schools(line)]
    if len(schools) < 2:
        return {}
    missing = [index for index, line in schools if not _pick_degree(line)]
    degree_rows = [(index, _pick_degree(line)) for index, line in enumerate(lines)
                   if not _real_schools(line) and _education_degree_line(line) and _pick_degree(line)]
    if not missing:
        return {}
    if not all(index > schools[-1][0] for index, _ in degree_rows):
        return {}
    if degree_rows and len(degree_rows) != len(missing):
        return {index: "" for index in missing}
    return dict(zip(missing, (degree for _, degree in degree_rows)))


def _merge_education_records(records: list[dict]) -> list[dict]:
    """Merge repeated renderings of one education record while preserving real multiple degrees."""
    merged: list[dict] = []
    for record in records:
        school_key = re.sub(r"\s+", "", str(record.get("school") or "")).lower()
        degree = normalize_degree(record.get("degree", ""))
        graduate_at = str(record.get("graduate_at") or "")
        existing = next((item for item in merged if (
            re.sub(r"\s+", "", str(item.get("school") or "")).lower() == school_key
            and (not degree or not item.get("degree") or normalize_degree(item.get("degree", "")) == degree)
            and (not graduate_at or not item.get("graduate_at") or item.get("graduate_at") == graduate_at)
        )), None)
        if existing is None:
            merged.append(record)
            continue
        for key in ("major", "degree", "graduate_at"):
            if not existing.get(key) and record.get(key):
                existing[key] = record[key]
    return merged


def _extract_education(text: str) -> list[dict]:
    """Associate nearby date/degree/major tokens with each school."""
    lines = _education_lines(text)
    column_degrees = _column_degrees(lines)
    records = []
    for index, line in enumerate(lines):
        schools = _real_schools(line)
        if not schools:
            continue
        context = _education_context(lines, index)
        line_dates = [_normalize_date(item) for item in DATE_RE.findall(line)]
        context_dates = [_normalize_date(item) for item in DATE_RE.findall(context)]
        context_degree = _pick_degree(context)
        context_major = _pick_major(context)
        for school in schools:
            school = _clean_school(school)
            if not school or school in FACULTY_ONLY_NAMES:
                continue
            record_degree = _pick_degree(line) or (column_degrees[index] if index in column_degrees else context_degree)
            record_degree = normalize_degree(record_degree) or _infer_degree_from_school(school) or record_degree
            record_major = (
                _pick_major(line)
                or _major_from_school_line(line, school)
                or context_major
            )
            records.append({
                "school": school,
                "major": record_major,
                "degree": record_degree,
                "graduate_at": (line_dates or context_dates)[-1] if (line_dates or context_dates) else "",
            })

    # A sparse master's/doctorate record must not be discarded just because
    # the same institution has a more complete undergraduate record.
    unique = []
    seen = set()
    for record in records:
        key = tuple(record.values())
        if key not in seen:
            seen.add(key)
            unique.append(record)
    return _merge_education_records(unique)


def _name_candidate_score(value: str, position: int = 0) -> tuple[int, str]:
    """Score a possible name and reject obvious job/department/file metadata."""
    candidate = _sanitize_name(re.sub(r"^(?:姓名|候选人|应聘者)\s*[：:]?", "", value))
    if not candidate or candidate.lower() in SECTION_HEADINGS:
        return -100, ""
    lowered = candidate.lower()
    if any(term.lower() in lowered for term in NAME_BLOCKLIST_TERMS):
        return -100, ""
    if any(token in candidate for token in ("大学", "学院", "学校", "公司", "集团", "部门", "中心")):
        return -100, ""
    if re.fullmatch(r"[\u4e00-\u9fff·]{2,5}", candidate):
        plain = candidate.replace("·", "")
        score = 5
        if plain and plain[0] in COMMON_CHINESE_SURNAMES:
            score += 8
        if 2 <= len(plain) <= 4:
            score += 3
        return score - min(position, 5), candidate
    if re.fullmatch(r"[A-Za-z][A-Za-z .'-]{1,39}", candidate):
        words = [item for item in candidate.split() if item]
        if 2 <= len(words) <= 4:
            return 6 - min(position, 3), candidate
    return -100, ""


def _name_from_filename(filename: str) -> str:
    if not filename:
        return ""
    stem = os.path.splitext(os.path.basename(filename))[0]
    stem = re.sub(r"\[(?:from)?[^\]]*\]", " ", stem, flags=re.I)
    parts = [item for item in re.split(r"[_\-—–|｜\s（）()\[\]]+", stem) if item]
    ranked = [_name_candidate_score(part, index) for index, part in enumerate(parts)]
    score, candidate = max(ranked, default=(-100, ""), key=lambda item: item[0])
    return candidate if score > 0 else ""


def _name_from_header(text: str) -> str:
    """Find an unlabeled display name near the top of a resume."""
    ranked = []
    for index, line in enumerate(line.strip() for line in text.splitlines() if line.strip()):
        if index >= 15:
            break
        candidate = re.sub(r"\s+", " ", line).strip(" ：:，,;；|")
        if len(candidate) > 40 or re.search(r"@|1[3-9]\d{9}", candidate):
            continue
        ranked.append(_name_candidate_score(candidate, index))
    score, candidate = max(ranked, default=(-100, ""), key=lambda item: item[0])
    # 无标签中文姓名必须具有姓氏证据；否则“手机”“最高学历”等短字段也会被误判。
    if re.fullmatch(r"[\u4e00-\u9fff·]{2,5}", candidate or ""):
        return candidate if score >= 12 else ""
    return candidate if score > 0 else ""


def _normalize_gender(value: str) -> str:
    value = value.strip().lower()
    if value in {"\u7537", "male", "m"}:
        return "\u7537"
    if value in {"\u5973", "female", "f"}:
        return "\u5973"
    return ""


def _extract_gender(text: str, filename: str = "") -> str:
    compact = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    for source in (text, compact):
        matched = GENDER_LABEL_RE.search(source)
        if matched:
            return _normalize_gender(matched.group(1))
    for line in text.splitlines():
        line = re.sub(r"\s+", " ", line).strip(" \uFF1A:，,;；|-")
        matched = re.fullmatch(
            r"(?:\u6027\u522b|gender|sex)\s*(?:[\uFF1A:]\s*)?(\u7537|\u5973|male|female|m|f)",
            line,
            re.I,
        )
        if matched:
            return _normalize_gender(matched.group(1))
    stem = os.path.splitext(os.path.basename(filename))[0]
    matched = re.search(r"(?:^|[_\-\s])([\u7537\u5973])(?:[_\-\s]|$)", stem)
    return matched.group(1) if matched else ""


def _extract_city(text: str) -> str:
    """提取带明确标签的现居城市，避免把学校或公司所在地误当作城市。"""
    patterns = (
        r"(?:现居住地|现居地|现居城市|所在城市|所在地|居住地|城市)\s*[：:|｜]\s*([^\r\n|，,;；]{1,30})",
        r"(?:Current\s+City|City|Location)\s*[:：|｜]\s*([^\r\n|,;]{2,40})",
    )
    for pattern in patterns:
        matched = re.search(pattern, text, flags=re.IGNORECASE)
        if not matched:
            continue
        value = re.split(
            r"(?:性别|年龄|手机(?:号码)?|电话|邮箱|电子邮箱|求职|应聘)\s*[：:]",
            matched.group(1), maxsplit=1,
        )[0]
        value = re.sub(r"\s+", " ", value).strip(" ：:，,;；|/-")
        chinese = re.match(r"[\u4e00-\u9fff]{2,12}", value)
        if chinese:
            return chinese.group(0)
        english = re.match(r"[A-Za-z][A-Za-z .'-]{1,39}", value)
        if english:
            return english.group(0).strip()
    return ""


def _is_section_heading(line: str) -> bool:
    normalized = re.sub(r"[\uFF1A:\s]+$", "", re.split(r"[|｜]", line, maxsplit=1)[0].strip()).lower()
    return normalized in ALL_SECTION_HEADINGS


def _extract_work_experience(text: str) -> list[dict]:
    return extract_work_experience(text)


def _sanitize_name(value: str) -> str:
    """Return only a likely person name, never a filename/section suffix."""
    value = re.sub(r"\s+", " ", str(value or "")).strip(" \uFF1A:，,;；|/")
    value = re.split(r"[_\-\[\]（）()（）]", value, maxsplit=1)[0].strip()
    value = re.sub(
        r"(?:\u4E2A\u4EBA)?\u7B80\u5386(?:\u6587\u4EF6|\u6587\u6863)?$|(?:resume|curriculum\s+vitae|cv)$",
        "",
        value,
        flags=re.IGNORECASE,
    ).strip(" \uFF1A:，,;；|/")
    chinese = re.match(r"[\u4e00-\u9fff·]{2,6}", value)
    if chinese:
        candidate = chinese.group(0).strip("·")
        return "" if candidate in SECTION_HEADINGS else candidate
    english = re.match(r"[A-Za-z][A-Za-z .'-]{1,39}", value)
    return english.group(0).strip() if english else ""


def parse_resume_fields(text: str, filename: str = "") -> dict:
    """Extract basic information, education and work experience from resume text."""
    fields = {
        "name": "", "phone": "", "email": "", "city": "", "age": None,
        "highest_education": "", "major": "", "education": [],
    }
    if not text:
        fields.update({"gender": "", "work_experience": []})
        return fields
    fields.update({"gender": "", "work_experience": []})
    text = _normalize_resume_text(text)

    compact = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    phone = re.search(r"1[3-9](?:[\s-]*\d){9}", text)
    if phone:
        fields["phone"] = re.sub(r"\D", "", phone.group(0))
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    if email:
        fields["email"] = email.group(0)
    fields["gender"] = _extract_gender(text, filename)
    fields["city"] = _extract_city(text)

    name_patterns = [
        r"(?:姓\s*名|名字|候选人|应聘者)\s*[：:|｜]\s*([^\r\n]{2,30})",
        r"(?:Name|Candidate\s+Name)\s*[:：|｜]\s*([^\r\n]{2,30})",
        r"(?:姓\s*名|名字)\s+([\u4e00-\u9fa5A-Za-z·]{2,20})",
    ]
    name_value = ""
    for pattern in name_patterns:
        matched = re.search(pattern, text, flags=re.IGNORECASE) or re.search(
            pattern, compact, flags=re.IGNORECASE,
        )
        if matched:
            name_value = matched.group(1)
            break
    if name_value:
        name_value = re.split(
            r"(?:性别|手机(?:号码)?|电话|邮箱|电子邮箱|城市|所在地|现居)\s*[：:]",
            name_value,
            maxsplit=1,
        )[0]
        name_value = re.sub(r"\s+", " ", name_value).strip(" ：:，,;；")
        name_match = re.match(r"[\u4e00-\u9fa5A-Za-z·][\u4e00-\u9fa5A-Za-z· .'-]{1,19}", name_value)
        if name_match:
            fields["name"] = name_match.group(0).strip()

    if not fields["name"]:
        fields["name"] = _name_from_header(text)

    if not fields["name"]:
        fields["name"] = _name_from_filename(filename)

    fields["name"] = _sanitize_name(fields["name"])
    fields["education"] = _extract_education(text)
    fields["age"] = _extract_age(text)
    fields["highest_education"] = highest_education_from_records(fields["education"], text)
    fields["major"] = major_from_records(fields["education"], text)
    fields["work_experience"] = _extract_work_experience(text)
    return fields


def _resume_quality(fields: dict, text: str, filename: str = "") -> tuple[int, list[str]]:
    """Estimate whether deterministic parsing is safe enough to use directly."""
    score = 0
    warnings = []
    if fields.get("name"):
        score += 22
        labelled = re.search(r"(?:姓名|名字|候选人|应聘者)\s*[：:|｜]", text, re.I)
        header_name = _name_from_header(text)
        if not labelled and not header_name and fields.get("name") == _name_from_filename(filename):
            score -= 8
            warnings.append("姓名仅来自文件名，建议核对")
    else:
        warnings.append("未识别到姓名")
    if fields.get("phone"):
        score += 15
    if fields.get("email"):
        score += 10
    if fields.get("city"):
        score += 5
    if fields.get("highest_education"):
        score += 10
    education = fields.get("education") or []
    if education:
        score += min(20, 8 + len(education) * 4)
        if any(not normalize_degree(record.get("degree", "")) for record in education):
            warnings.append("部分教育经历未识别到明确学历，请核对本科/硕士/博士")
        if len(education) > 4:
            warnings.append("教育经历数量异常，建议核对")
    elif re.search(r"教育经历|教育背景|学历背景|education", text, re.I):
        warnings.append("存在教育章节但未提取到院校")
    education_degrees = {_pick_degree(line) for line in _education_lines(text) if _education_degree_line(line)} - {""}
    recorded_degrees = {normalize_degree(record.get("degree", "")) for record in education}
    if education and education_degrees - recorded_degrees:
        warnings.append("教育章节存在尚未关联的学历，请核对各段教育经历")
    work = fields.get("work_experience") or []
    if work:
        score += min(18, 8 + len(work) * 3)
        if any(not record.get("start") or not record.get("end") for record in work):
            warnings.append("部分工作经历未识别到完整任职时间，请核对")
    elif re.search(r"工作经历|工作经验|职业经历|work experience", text, re.I):
        warnings.append("存在工作章节但未提取到公司")
    return max(0, min(score, 100)), warnings


def _valid_ai_name(value: object) -> str:
    score, candidate = _name_candidate_score(str(value or ""))
    return candidate if score > 0 else ""


def _clean_ai_education(value: object) -> list[dict]:
    if not isinstance(value, list):
        return []
    records = []
    for raw in value[:10]:
        if not isinstance(raw, dict):
            continue
        school = _clean_school(str(raw.get("school") or ""))
        if not school or school in FACULTY_ONLY_NAMES:
            continue
        records.append({
            "school": school[:80],
            "major": str(raw.get("major") or "").strip()[:80],
            "degree": normalize_degree(raw.get("degree", "")),
            "graduate_at": _normalize_date(str(raw.get("graduate_at") or ""))[:20],
        })
    return _merge_education_records(records)


def _clean_ai_work(value: object) -> list[dict]:
    if not isinstance(value, list):
        return []
    records = []
    for raw in value[:15]:
        if not isinstance(raw, dict):
            continue
        company = str(raw.get("company") or "").strip()
        if not company:
            continue
        records.append({
            "company": company[:120],
            "position": str(raw.get("position") or "").strip()[:80],
            "start": _normalize_work_date(str(raw.get("start") or ""))[:20],
            "end": _normalize_work_date(str(raw.get("end") or ""))[:20],
            "desc": str(raw.get("desc") or "").strip()[:500],
        })
    return records


def _merge_education_evidence(records: list[dict], ai_records: list[dict], text: str) -> list[dict]:
    """Supplement missing histories without promoting a degree on AI say-so."""
    source_records = _extract_education(text)
    lines = _education_lines(text)
    merged = [dict(record) for record in records]

    def key(value):
        return re.sub(r"\s+", "", _normalize_resume_text(value)).casefold()

    def same_period(left, right):
        a, b = left.get("graduate_at", ""), right.get("graduate_at", "")
        return not a or not b or a == b or a[:4] == b[:4] and (len(a) == 4 or len(b) == 4)

    for incoming in ai_records:
        school = key(incoming["school"])
        source = [record for record in source_records
                  if key(record["school"]) == school and same_period(record, incoming)]
        degree = normalize_degree(incoming.get("degree", ""))
        # Recheck both institution and degree in the education section, not
        # in a work-project description or a copied job requirement.
        indexes = [index for index, line in enumerate(lines) if school in key(line) and _education_degree_line(line)]
        if not indexes:
            continue
        context = "\n".join(_education_context(lines, index) for index in indexes)
        supported = any(normalize_degree(record.get("degree", "")) == degree for record in source) if source else (
            degree in {_pick_degree(line) for line in context.splitlines() if _education_degree_line(line)}
        )
        if degree and not supported:
            continue
        record = dict(incoming)
        existing = next((item for item in merged if key(item["school"]) == school
                         and same_period(item, record)
                         and (not item.get("degree") or normalize_degree(item["degree"]) == degree)), None)
        if existing is None and degree:
            # Repair a wrong degree only when the same dated school record has
            # explicit source evidence; never overwrite a distinct earlier degree.
            existing = next((item for item in merged if key(item["school"]) == school
                             and item.get("graduate_at") and record.get("graduate_at")
                             and same_period(item, record)
                             and source and not any(normalize_degree(origin.get("degree", "")) == normalize_degree(item.get("degree", ""))
                                                    for origin in source)), None)
            if existing is not None:
                existing["degree"] = degree
        if existing is not None:
            for field in ("degree", "major", "graduate_at"):
                if not existing.get(field) and record.get(field):
                    existing[field] = record[field]
        else:
            merged.append(record)
    return _merge_education_records(merged)


def _merge_ai_fields(fields: dict, ai_fields: dict, text: str, filename: str) -> dict:
    """Use AI only for missing or explicitly low-confidence deterministic fields."""
    result = dict(fields)
    labelled_name = bool(re.search(r"(?:姓名|名字|候选人|应聘者)\s*[：:|｜]", text, re.I))
    ai_name = _valid_ai_name(ai_fields.get("name"))
    if ai_name and (not result.get("name") or (
        not labelled_name and result.get("name") == _name_from_filename(filename)
    )):
        result["name"] = ai_name

    validators = {
        "phone": lambda value: re.sub(r"\D", "", str(value or "")) if re.fullmatch(
            r"1[3-9](?:[\s-]*\d){9}", str(value or "").strip(),
        ) else "",
        "email": lambda value: str(value or "").strip() if re.fullmatch(
            r"[\w.+-]+@[\w-]+\.[\w.]+", str(value or "").strip(),
        ) else "",
        "city": lambda value: str(value or "").strip()[:30],
        "gender": lambda value: _normalize_gender(str(value or "")),
    }
    for key, validator in validators.items():
        if not result.get(key):
            cleaned = validator(ai_fields.get(key))
            if cleaned:
                result[key] = cleaned

    if result.get("age") is None:
        try:
            age = int(ai_fields.get("age"))
            if 16 <= age <= 75:
                result["age"] = age
        except (TypeError, ValueError):
            pass

    evidence_text = re.sub(r"\s+", "", _normalize_resume_text(text)).lower()

    def has_evidence(value: str) -> bool:
        value = re.sub(r"\s+", "", _normalize_resume_text(value)).lower()
        return bool(value and value in evidence_text)

    ai_education = [
        record for record in _clean_ai_education(ai_fields.get("education"))
        if has_evidence(record["school"])
    ]
    if ai_education:
        result["education"] = _merge_education_evidence(result.get("education") or [], ai_education, text)
    ai_work = [
        record for record in _clean_ai_work(ai_fields.get("work_experience"))
        if has_evidence(record["company"])
    ]
    if ai_work and not result.get("work_experience"):
        result["work_experience"] = ai_work

    if result.get("education"):
        result["highest_education"] = highest_education_from_records(result["education"], text)
        result["major"] = major_from_records(result["education"])
    else:
        ai_degree = normalize_degree(ai_fields.get("highest_education", ""))
        if ai_degree and has_evidence(ai_degree) and not result.get("highest_education"):
            result["highest_education"] = ai_degree
        if not result.get("major") and has_evidence(str(ai_fields.get("major") or "")):
            result["major"] = str(ai_fields.get("major") or "").strip()[:80]
    return result


def parse_resume_file_with_text(file_path: str, original_filename: str = ""):
    """Return parsed fields, status and source text for downstream matching."""
    ext = os.path.splitext(file_path)[1].lower()
    empty = {
        "name": "", "phone": "", "email": "", "city": "", "age": None,
        "highest_education": "", "major": "", "education": [],
    }
    empty.update({"gender": "", "work_experience": []})
    text, method, ocr_confidence = extract_text_with_method(file_path)
    if not text or not text.strip():
        empty["_parse_meta"] = {
            "method": method, "confidence": 0,
            "warnings": ["未能从文件中提取文字，请人工录入"],
        }
        return empty, "failed", ""
    fields = parse_resume_fields(text, original_filename)
    confidence, warnings = _resume_quality(fields, text, original_filename)
    ai_applied = False
    try:
        from common.resume_ai import extract_resume_fields, is_configured

        if is_configured() and (confidence < 80 or warnings):
            ai_fields, ai_error = extract_resume_fields(text, original_filename)
            if ai_fields:
                fields = _merge_ai_fields(fields, ai_fields, text, original_filename)
                ai_applied = True
                confidence, warnings = _resume_quality(fields, text, original_filename)
            elif ai_error:
                warnings.append(ai_error)
    except Exception:
        logger.exception("AI 简历兜底处理失败")
        warnings.append("AI 兜底处理失败，已保留规则解析结果")
    if method == "ocr" and ocr_confidence:
        confidence = min(confidence, round(ocr_confidence * 100))
        if ocr_confidence < 0.75:
            warnings.append("OCR 文字置信度偏低，请重点核对")
    fields["_parse_meta"] = {
        "method": f"{method}+ai" if ai_applied else method,
        "confidence": confidence,
        "warnings": list(dict.fromkeys(warnings)),
    }
    if not any([
        fields["name"], fields["gender"], fields["phone"], fields["email"],
        fields["education"], fields["work_experience"],
    ]):
        return fields, "failed", text[:100000]
    return fields, "system", text[:100000]


def parse_resume_file(file_path: str, original_filename: str = ""):
    """Return (fields, parse_status): system=parsed, failed=manual review."""
    fields, status, _ = parse_resume_file_with_text(file_path, original_filename)
    return fields, status
